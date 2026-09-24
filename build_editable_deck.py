import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def build_editable_deck():
    prs = Presentation()
    # 16:9 Widescreen format
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    slide = prs.slides.add_slide(blank_layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    
    # -------------------------------------------------------------------------
    # 1. TOP HEADER
    # -------------------------------------------------------------------------
    # Left Brand Logo Text
    brand_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.18), Inches(3.2), Inches(0.65))
    tf = brand_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "CIVIC CONNECT"
    p.font.name = "Arial Black"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 44, 89)
    p2 = tf.add_paragraph()
    p2.text = "Govt. of Jharkhand Civic Innovation"
    p2.font.name = "Arial"
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = RGBColor(100, 116, 139)

    # Center Main Title
    title_box = slide.shapes.add_textbox(Inches(3.8), Inches(0.15), Inches(5.8), Inches(0.7))
    tf_t = title_box.text_frame
    p_t = tf_t.paragraphs[0]
    p_t.text = "TECHNICAL APPROACH"
    p_t.alignment = PP_ALIGN.CENTER
    p_t.font.name = "Arial Black"
    p_t.font.size = Pt(24)
    p_t.font.bold = True
    p_t.font.color.rgb = RGBColor(15, 23, 42)

    # Right Hackathon Badge
    sih_box = slide.shapes.add_textbox(Inches(10.2), Inches(0.15), Inches(2.7), Inches(0.7))
    tf_s = sih_box.text_frame
    p_s = tf_s.paragraphs[0]
    p_s.text = "SMART INDIA HACKATHON 2025"
    p_s.alignment = PP_ALIGN.RIGHT
    p_s.font.name = "Arial"
    p_s.font.size = Pt(10.5)
    p_s.font.bold = True
    p_s.font.color.rgb = RGBColor(225, 112, 85)
    p_s2 = tf_s.add_paragraph()
    p_s2.text = "Idea Submission | Slide 3"
    p_s2.alignment = PP_ALIGN.RIGHT
    p_s2.font.size = Pt(8.5)
    p_s2.font.color.rgb = RGBColor(100, 116, 139)

    # -------------------------------------------------------------------------
    # 2. LEFT COLUMN: 4 SPECIFICATION CARDS (Pure Software Stack)
    # -------------------------------------------------------------------------
    left_x = Inches(0.4)
    left_w = Inches(2.7)

    specs = [
        ("System Environment", "• Java 17 & Python 3.10+\n• React 18 & Vite Client\n• Spring Boot 3.4 (Port 5000)\n• FastAPI AI (Port 8000)", Inches(0.95), Inches(1.15), RGBColor(59, 130, 246)),
        ("Frameworks & Libraries", "• Spring Boot 3.4 & Security\n• FastAPI & PyTorch\n• Sentence-Transformers\n• Pillow (ELA Forensics)", Inches(2.25), Inches(1.15), RGBColor(79, 70, 229)),
        ("Storage Technologies", "• MongoDB 7.0 Database\n• Local File Evidence Store\n• Vector Similarity Cache\n• Browser LocalStorage", Inches(3.55), Inches(1.05), RGBColor(14, 165, 233)),
        ("Integration & Protocols", "• MCP Server (72 Tools)\n• RESTful JSON HTTP APIs\n• JWT Token Authentication\n• Role-Based Access (RBAC)", Inches(4.75), Inches(1.05), RGBColor(2, 132, 199))
    ]

    for title, text_content, y_pos, card_h, pill_col in specs:
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, y_pos, left_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(248, 250, 252)
        card.line.color.rgb = RGBColor(203, 213, 225)
        card.line.width = Pt(1)

        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x + Inches(0.12), y_pos - Inches(0.1), left_w - Inches(0.24), Inches(0.26))
        pill.fill.solid()
        pill.fill.fore_color.rgb = pill_col
        pill.line.fill.background()
        p = pill.text_frame.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Arial"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.2)
        tf.margin_left = Inches(0.1)
        tf.margin_right = Inches(0.1)
        p = tf.paragraphs[0]
        p.text = text_content
        p.font.name = "Arial"
        p.font.size = Pt(7.8)
        p.font.color.rgb = RGBColor(30, 41, 59)

    # Bottom Left Yellow Box (Scope, Design, Dev, Deploy, Maintain)
    yellow_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, Inches(5.95), Inches(4.3), Inches(1.18))
    yellow_box.fill.solid()
    yellow_box.fill.fore_color.rgb = RGBColor(254, 240, 138)
    yellow_box.line.color.rgb = RGBColor(234, 179, 8)
    yellow_box.line.width = Pt(1.5)

    tf_y = yellow_box.text_frame
    tf_y.word_wrap = True
    tf_y.margin_left = Inches(0.12)
    tf_y.margin_top = Inches(0.06)
    
    phases = [
        "• Scope: Ingest civic complaints, match R&D with CSR funds, track SLAs",
        "• Design: 3-tier microservice architecture (React + Spring Boot + FastAPI + MCP)",
        "• Development: Sentence-Transformers NLP, ELA photo authenticity & MCDA matching",
        "• Deployment: Containerized microservices with MongoDB & JWT security",
        "• Maintenance: Real-time SLA tracking, multi-district analytics & audit logs"
    ]
    for i, phase in enumerate(phases):
        p = tf_y.paragraphs[0] if i == 0 else tf_y.add_paragraph()
        p.text = phase
        p.font.name = "Arial"
        p.font.size = Pt(7.3)
        p.font.color.rgb = RGBColor(113, 63, 18)
        if i == 0:
            p.font.bold = True

    # -------------------------------------------------------------------------
    # 3. CENTER PIPELINE: 6 INTERCONNECTED STAGES (A -> B -> C -> D -> E -> F)
    # -------------------------------------------------------------------------
    
    # Column A: Stakeholder Portals
    colA_x = Inches(3.25)
    colA_w = Inches(1.25)
    
    cardA = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, colA_x, Inches(0.95), colA_w, Inches(4.85))
    cardA.fill.solid()
    cardA.fill.fore_color.rgb = RGBColor(241, 245, 249)
    cardA.line.color.rgb = RGBColor(203, 213, 225)
    
    hdrA = slide.shapes.add_shape(MSO_SHAPE.OVAL, colA_x + Inches(0.42), Inches(1.02), Inches(0.4), Inches(0.4))
    hdrA.fill.solid()
    hdrA.fill.fore_color.rgb = RGBColor(30, 41, 59)
    hdrA.line.fill.background()
    p = hdrA.text_frame.paragraphs[0]
    p.text = "A"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    lblA = slide.shapes.add_textbox(colA_x, Inches(1.45), colA_w, Inches(0.3))
    p = lblA.text_frame.paragraphs[0]
    p.text = "Stakeholder\nPortals"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 23, 42)
    
    portals = [
        ("📱 Citizen", "Web/Mobile"),
        ("🖥️ Admin", "Command"),
        ("🎓 Univ.", "R&D Hub"),
        ("🏭 Industry", "CSR Portal")
    ]
    for i, (pname, psub) in enumerate(portals):
        sub_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, colA_x + Inches(0.08), Inches(1.95 + i*0.88), colA_w - Inches(0.16), Inches(0.78))
        sub_card.fill.solid()
        sub_card.fill.fore_color.rgb = RGBColor(255, 255, 255)
        sub_card.line.color.rgb = RGBColor(226, 232, 240)
        tf = sub_card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.08)
        p = tf.paragraphs[0]
        p.text = pname
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = RGBColor(30, 41, 59)
        p2 = tf.add_paragraph()
        p2.text = psub
        p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(7.2)
        p2.font.color.rgb = RGBColor(100, 116, 139)

    # Column B: Input Box
    colB_x = Inches(4.6)
    colB_w = Inches(1.25)
    
    cardB = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, colB_x, Inches(0.95), colB_w, Inches(4.85))
    cardB.fill.solid()
    cardB.fill.fore_color.rgb = RGBColor(254, 243, 199)
    cardB.line.color.rgb = RGBColor(252, 211, 77)
    
    hdrB = slide.shapes.add_shape(MSO_SHAPE.OVAL, colB_x + Inches(0.42), Inches(1.02), Inches(0.4), Inches(0.4))
    hdrB.fill.solid()
    hdrB.fill.fore_color.rgb = RGBColor(180, 83, 9)
    hdrB.line.fill.background()
    p = hdrB.text_frame.paragraphs[0]
    p.text = "B"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    lblB = slide.shapes.add_textbox(colB_x, Inches(1.45), colB_w, Inches(0.3))
    p = lblB.text_frame.paragraphs[0]
    p.text = "Input\nData"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(113, 63, 18)
    
    inputs = [
        ("📝 Text", "['Hindi', 'English'\nGrievances]"),
        ("📷 Photo", "['.jpg', '.png'\nGeo Evidence]"),
        ("📄 Docs", "['.pdf' R&D / CSR\nProposals]")
    ]
    for i, (iname, isub) in enumerate(inputs):
        sub_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, colB_x + Inches(0.08), Inches(2.05 + i*1.15), colB_w - Inches(0.16), Inches(1.05))
        sub_card.fill.solid()
        sub_card.fill.fore_color.rgb = RGBColor(255, 255, 255)
        sub_card.line.color.rgb = RGBColor(245, 158, 11)
        tf = sub_card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.1)
        p = tf.paragraphs[0]
        p.text = iname
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(8.8)
        p.font.bold = True
        p.font.color.rgb = RGBColor(146, 64, 14)
        p2 = tf.add_paragraph()
        p2.text = isub
        p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(7.2)
        p2.font.color.rgb = RGBColor(71, 85, 105)

    # Column C: Data Cleansing
    colC_x = Inches(5.95)
    colC_w = Inches(1.3)
    
    cardC = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, colC_x, Inches(0.95), colC_w, Inches(4.85))
    cardC.fill.solid()
    cardC.fill.fore_color.rgb = RGBColor(240, 253, 250)
    cardC.line.color.rgb = RGBColor(153, 246, 228)
    
    hdrC = slide.shapes.add_shape(MSO_SHAPE.OVAL, colC_x + Inches(0.45), Inches(1.02), Inches(0.4), Inches(0.4))
    hdrC.fill.solid()
    hdrC.fill.fore_color.rgb = RGBColor(13, 148, 136)
    hdrC.line.fill.background()
    p = hdrC.text_frame.paragraphs[0]
    p.text = "C"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    lblC = slide.shapes.add_textbox(colC_x, Inches(1.45), colC_w, Inches(0.3))
    p = lblC.text_frame.paragraphs[0]
    p.text = "Data\nCleansing"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(17, 94, 89)
    
    cleans = [
        ("🔍 ELA Forensics", "Error Level Analysis\n(Fake photo check)"),
        ("📑 Sentence Dedup", "MiniLM embeddings\nmerge duplicates"),
        ("🛡️ Sanitization", "NoSQL / XSS\nsecurity filters")
    ]
    for i, (cname, csub) in enumerate(cleans):
        sub_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, colC_x + Inches(0.08), Inches(2.05 + i*1.15), colC_w - Inches(0.16), Inches(1.05))
        sub_card.fill.solid()
        sub_card.fill.fore_color.rgb = RGBColor(255, 255, 255)
        sub_card.line.color.rgb = RGBColor(204, 251, 241)
        tf = sub_card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.1)
        p = tf.paragraphs[0]
        p.text = cname
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(8.2)
        p.font.bold = True
        p.font.color.rgb = RGBColor(13, 148, 136)
        p2 = tf.add_paragraph()
        p2.text = csub
        p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(7.0)
        p2.font.color.rgb = RGBColor(71, 85, 105)

    # Column D: Gateway & Pre-filter
    colD_x = Inches(7.35)
    colD_w = Inches(1.2)
    
    cardD = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, colD_x, Inches(0.95), colD_w, Inches(4.85))
    cardD.fill.solid()
    cardD.fill.fore_color.rgb = RGBColor(241, 245, 249)
    cardD.line.color.rgb = RGBColor(203, 213, 225)
    
    hdrD = slide.shapes.add_shape(MSO_SHAPE.OVAL, colD_x + Inches(0.4), Inches(1.02), Inches(0.4), Inches(0.4))
    hdrD.fill.solid()
    hdrD.fill.fore_color.rgb = RGBColor(30, 41, 59)
    hdrD.line.fill.background()
    p = hdrD.text_frame.paragraphs[0]
    p.text = "D"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    lblD = slide.shapes.add_textbox(colD_x, Inches(1.45), colD_w, Inches(0.3))
    p = lblD.text_frame.paragraphs[0]
    p.text = "Gateway &\nPre-Filter"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 23, 42)
    
    gateways = [
        ("🍃 Spring Boot", "REST API Gateway\n(Port 5000)"),
        ("⚡ Urgency Filter", "Anomaly & SLA\nurgency sorter"),
        ("🔐 Role Routing", "Citizen, Admin,\nUniv, Industry")
    ]
    for i, (gname, gsub) in enumerate(gateways):
        sub_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, colD_x + Inches(0.08), Inches(2.05 + i*1.15), colD_w - Inches(0.16), Inches(1.05))
        sub_card.fill.solid()
        sub_card.fill.fore_color.rgb = RGBColor(255, 255, 255)
        sub_card.line.color.rgb = RGBColor(226, 232, 240)
        tf = sub_card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.1)
        p = tf.paragraphs[0]
        p.text = gname
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(8.2)
        p.font.bold = True
        p.font.color.rgb = RGBColor(30, 41, 59)
        p2 = tf.add_paragraph()
        p2.text = gsub
        p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(7.0)
        p2.font.color.rgb = RGBColor(100, 116, 139)

    # Column E: AI/ML & MCP Services (Core Engine)
    colE_x = Inches(8.65)
    colE_w = Inches(1.9)
    
    cardE = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, colE_x, Inches(0.95), colE_w, Inches(4.85))
    cardE.fill.solid()
    cardE.fill.fore_color.rgb = RGBColor(238, 242, 255)
    cardE.line.color.rgb = RGBColor(199, 210, 254)
    cardE.line.width = Pt(1.5)
    
    hdrE = slide.shapes.add_shape(MSO_SHAPE.OVAL, colE_x + Inches(0.75), Inches(1.02), Inches(0.4), Inches(0.4))
    hdrE.fill.solid()
    hdrE.fill.fore_color.rgb = RGBColor(67, 56, 202)
    hdrE.line.fill.background()
    p = hdrE.text_frame.paragraphs[0]
    p.text = "E"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    lblE = slide.shapes.add_textbox(colE_x, Inches(1.45), colE_w, Inches(0.3))
    p = lblE.text_frame.paragraphs[0]
    p.text = "AI / ML & MCP Services"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(9.0)
    p.font.bold = True
    p.font.color.rgb = RGBColor(49, 46, 129)
    
    ai_services = [
        ("🤖 Sentence-Transformers", "Multilingual NLP & Hindi/Eng\ndomain categorization"),
        ("👁️ ELA Classifier (PyTorch)", "Authenticity classifier for\nphysical evidence validation"),
        ("🤝 Univ-CSR Matchmaker", "MCDA algorithm matching\nR&D labs with CSR grants"),
        ("⚡ 72 MCP Decision Tools", "Gate 1, 2, 3 verification &\nMilestone decision support")
    ]
    for i, (aname, asub) in enumerate(ai_services):
        sub_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, colE_x + Inches(0.08), Inches(1.85 + i*0.95), colE_w - Inches(0.16), Inches(0.85))
        sub_card.fill.solid()
        sub_card.fill.fore_color.rgb = RGBColor(255, 255, 255)
        sub_card.line.color.rgb = RGBColor(224, 231, 255)
        tf = sub_card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.06)
        p = tf.paragraphs[0]
        p.text = aname
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(8.0)
        p.font.bold = True
        p.font.color.rgb = RGBColor(67, 56, 202)
        p2 = tf.add_paragraph()
        p2.text = asub
        p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(6.8)
        p2.font.color.rgb = RGBColor(71, 85, 105)

    # Database label inside E at bottom
    db_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, colE_x + Inches(0.15), Inches(5.38), colE_w - Inches(0.3), Inches(0.32))
    db_box.fill.solid()
    db_box.fill.fore_color.rgb = RGBColor(22, 163, 74)
    db_box.line.fill.background()
    p = db_box.text_frame.paragraphs[0]
    p.text = "🍃 MongoDB 7.0 Database"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(8.0)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    # Column F: Dashboards & Security
    colF_x = Inches(10.65)
    colF_w = Inches(2.25)

    # Top Right Product Status Sticky Note
    status_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, colF_x, Inches(0.95), colF_w, Inches(1.6))
    status_card.fill.solid()
    status_card.fill.fore_color.rgb = RGBColor(254, 243, 199)
    status_card.line.color.rgb = RGBColor(245, 158, 11)
    status_card.line.width = Pt(1.5)

    pin = slide.shapes.add_shape(MSO_SHAPE.OVAL, colF_x + colF_w - Inches(0.35), Inches(0.85), Inches(0.22), Inches(0.22))
    pin.fill.solid()
    pin.fill.fore_color.rgb = RGBColor(220, 38, 38)
    pin.line.fill.background()

    tf_s = status_card.text_frame
    tf_s.word_wrap = True
    tf_s.margin_top = Inches(0.08)
    p = tf_s.paragraphs[0]
    p.text = "Product Status"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(146, 64, 14)
    p2 = tf_s.add_paragraph()
    p2.text = "100% Core System Built & Verified!\n(54/54 Tests Passed, Production Ready)"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(8.2)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(30, 64, 175)

    # Middle Right Auth & Security Box
    sec_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, colF_x, Inches(2.65), colF_w, Inches(1.95))
    sec_box.fill.solid()
    sec_box.fill.fore_color.rgb = RGBColor(245, 243, 255)
    sec_box.line.color.rgb = RGBColor(221, 214, 254)
    sec_box.line.width = Pt(1)

    tf_sec = sec_box.text_frame
    tf_sec.word_wrap = True
    tf_sec.margin_top = Inches(0.08)
    p = tf_sec.paragraphs[0]
    p.text = "🔒 Security & Governance"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(109, 40, 217)
    p2 = tf_sec.add_paragraph()
    p2.text = "• JWT + BCrypt Auth\n• Role-Based Access (RBAC)\n• AES-256 Data Encryption\n• SHA-256 Tamper-Proof Audit\n• Human-In-The-Loop AI Gates"
    p2.font.size = Pt(7.4)
    p2.font.color.rgb = RGBColor(68, 64, 60)

    # Bottom Right TRL, DRL, IRL Badges
    badges = [
        ("7", "TECHNOLOGY\nREADINESS LEVEL", RGBColor(14, 116, 144), Inches(4.75)),
        ("6", "DEPLOYMENT\nREADINESS LEVEL", RGBColor(161, 98, 7), Inches(5.45)),
        ("5", "INVESTMENT\nREADINESS LEVEL", RGBColor(30, 58, 138), Inches(6.15))
    ]
    for num, label, col, y_pos in badges:
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, colF_x + Inches(0.1), y_pos, Inches(0.48), Inches(0.48))
        circ.fill.solid()
        circ.fill.fore_color.rgb = col
        circ.line.fill.background()
        p = circ.text_frame.paragraphs[0]
        p.text = num
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)

        lbl = slide.shapes.add_textbox(colF_x + Inches(0.65), y_pos - Inches(0.05), Inches(1.5), Inches(0.55))
        tf = lbl.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = label
        p.font.size = Pt(6.8)
        p.font.bold = True
        p.font.color.rgb = col

    # Footer
    footer = slide.shapes.add_textbox(Inches(4.5), Inches(7.15), Inches(4.4), Inches(0.3))
    p = footer.text_frame.paragraphs[0]
    p.text = "@SIH Idea submission - CIVIC CONNECT"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = RGBColor(100, 116, 139)

    pg = slide.shapes.add_textbox(Inches(12.2), Inches(7.15), Inches(0.8), Inches(0.3))
    p = pg.text_frame.paragraphs[0]
    p.text = "3"
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = RGBColor(71, 85, 105)

    # -------------------------------------------------------------------------
    # SLIDE 2: VERIFIED TECHNOLOGY STACK
    # -------------------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    slide2.background.fill.solid()
    slide2.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    border2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(0.3), Inches(12.533), Inches(6.9))
    border2.fill.background()
    border2.line.color.rgb = RGBColor(203, 213, 225)
    border2.line.width = Pt(1.5)

    t2 = slide2.shapes.add_textbox(Inches(3.5), Inches(0.45), Inches(6.333), Inches(0.7))
    p = t2.text_frame.paragraphs[0]
    p.text = "Technology Stack"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Georgia"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.underline = True
    p.font.color.rgb = RGBColor(24, 43, 73)

    tech_data = [
        ("Frontend", [("React 18", RGBColor(6, 182, 212)), ("Vite", RGBColor(168, 85, 247)), ("Bootstrap 5", RGBColor(147, 51, 234)), ("JavaScript", RGBColor(234, 179, 8)), ("React Router", RGBColor(239, 68, 68)), ("Lucide Icons", RGBColor(59, 130, 246))], Inches(1.3)),
        ("Backend", [("Java 17", RGBColor(234, 88, 12)), ("Spring Boot 3.4", RGBColor(22, 163, 74)), ("Spring Security", RGBColor(34, 197, 94)), ("JWT (JJWT)", RGBColor(219, 39, 119)), ("Maven", RGBColor(185, 28, 28))], Inches(2.7)),
        ("AI / ML Microservice", [("Python 3", RGBColor(30, 64, 175)), ("FastAPI", RGBColor(13, 148, 136)), ("PyTorch", RGBColor(234, 88, 12)), ("Transformers", RGBColor(217, 119, 6)), ("Sentence-Transformers", RGBColor(249, 115, 22)), ("Pillow (ELA)", RGBColor(220, 38, 38)), ("Scikit-Learn", RGBColor(234, 88, 12))], Inches(4.1)),
        ("Database & Engine", [("MongoDB 7.0", RGBColor(22, 163, 74)), ("MCP Server (72 Tools)", RGBColor(79, 70, 229)), ("Uvicorn", RGBColor(13, 148, 136))], Inches(5.5))
    ]

    for label, items, y_pos in tech_data:
        lbl_box = slide2.shapes.add_textbox(Inches(0.6), y_pos, Inches(2.2), Inches(1.0))
        p = lbl_box.text_frame.paragraphs[0]
        p.text = label
        p.font.name = "Arial"
        p.font.size = Pt(12.5)
        p.font.bold = True
        p.font.color.rgb = RGBColor(15, 23, 42)

        x_start = Inches(2.9)
        spacing = Inches(1.55)
        for i, (name, col) in enumerate(items):
            bx = x_start + i * spacing
            card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, y_pos, Inches(1.42), Inches(0.95))
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(248, 250, 252)
            card.line.color.rgb = RGBColor(226, 232, 240)
            card.line.width = Pt(1)

            tf_b = card.text_frame
            tf_b.word_wrap = True
            tf_b.margin_top = Inches(0.12)
            p = tf_b.paragraphs[0]
            p.text = name
            p.alignment = PP_ALIGN.CENTER
            p.font.name = "Arial"
            p.font.size = Pt(9.0)
            p.font.bold = True
            p.font.color.rgb = col

    output_file = "CivicConnect_Technical_Approach_Architecture.pptx"
    prs.save(output_file)
    print(f"Presentation generated successfully: {os.path.abspath(output_file)}")

if __name__ == "__main__":
    build_editable_deck()
