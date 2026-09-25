import os
import sys
import subprocess
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_html_version(output_html_path):
    html_content = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<style>
  * { box-sizing: border-box; margin: 0; padding: 0; font-family: 'Segoe UI', -apple-system, BlinkMacSystemFont, Roboto, Helvetica, Arial, sans-serif; }
  body {
    width: 1920px;
    height: 1080px;
    background: #F1F5F9;
    padding: 24px 32px;
    display: flex;
    flex-direction: column;
    justify-content: space-between;
    overflow: hidden;
  }
  
  /* Top Banner */
  .header-banner {
    background: linear-gradient(135deg, #0A192F 0%, #1E3A8A 100%);
    border-radius: 14px;
    padding: 16px 24px;
    text-align: center;
    color: #FFFFFF;
    box-shadow: 0 4px 14px rgba(0, 0, 0, 0.15);
    margin-bottom: 16px;
  }
  .header-title {
    font-size: 32px;
    font-weight: 900;
    letter-spacing: 2px;
    margin-bottom: 4px;
    text-transform: uppercase;
  }
  .header-subtitle {
    font-size: 18px;
    font-weight: 700;
    color: #93C5FD;
    margin-bottom: 4px;
  }
  .header-tags {
    font-size: 13.5px;
    color: #E2E8F0;
    font-weight: 500;
    letter-spacing: 0.5px;
  }
  
  /* Grid Rows */
  .row-3-col {
    display: grid;
    grid-template-columns: 1fr 1fr 1fr;
    gap: 20px;
    margin-bottom: 14px;
  }
  
  .row-wide {
    margin-bottom: 14px;
  }
  
  /* Card Styling */
  .card {
    border-radius: 12px;
    border-width: 1.5px;
    border-style: solid;
    background: #FFFFFF;
    box-shadow: 0 2px 8px rgba(0,0,0,0.06);
    display: flex;
    flex-direction: column;
    overflow: hidden;
  }
  
  .card-header {
    display: flex;
    align-items: center;
    gap: 12px;
    padding: 10px 16px;
    border-bottom: 1px solid rgba(0,0,0,0.06);
  }
  
  .badge {
    width: 28px;
    height: 28px;
    border-radius: 50%;
    color: #FFFFFF;
    font-weight: 900;
    font-size: 15px;
    display: flex;
    align-items: center;
    justify-content: center;
    flex-shrink: 0;
    box-shadow: 0 2px 4px rgba(0,0,0,0.15);
  }
  
  .card-title {
    font-size: 16px;
    font-weight: 800;
    letter-spacing: 0.5px;
    text-transform: uppercase;
  }
  
  .card-body {
    padding: 12px 18px;
    display: flex;
    flex-direction: column;
    gap: 6px;
    flex: 1;
    justify-content: center;
  }
  
  .bullet-item {
    font-size: 13.5px;
    color: #1E293B;
    display: flex;
    align-items: flex-start;
    gap: 8px;
    line-height: 1.35;
    font-weight: 500;
  }
  .bullet-dot {
    color: #64748B;
    font-weight: 900;
    font-size: 14px;
    line-height: 1;
    margin-top: 1px;
  }
  
  /* Colors */
  /* 1. Yellow/Amber */
  .card-1 { border-color: #FCD34D; background: #FFFBEB; }
  .card-1 .card-header { background: #FEF3C7; }
  .card-1 .badge { background: #D97706; }
  .card-1 .card-title { color: #92400E; }
  
  /* 2. Green */
  .card-2 { border-color: #A7F3D0; background: #ECFDF5; }
  .card-2 .card-header { background: #D1FAE5; }
  .card-2 .badge { background: #059669; }
  .card-2 .card-title { color: #065F46; }
  
  /* 3. Orange */
  .card-3 { border-color: #FDBA74; background: #FFF7ED; }
  .card-3 .card-header { background: #FFEDD5; }
  .card-3 .badge { background: #EA580C; }
  .card-3 .card-title { color: #9A3412; }
  
  /* 4. Blue Wide */
  .card-4 { border-color: #93C5FD; background: #EFF6FF; }
  .card-4 .card-header { background: #DBEAFE; }
  .card-4 .badge { background: #2563EB; }
  .card-4 .card-title { color: #1E40AF; }
  
  .wide-grid {
    display: grid;
    grid-template-columns: repeat(5, 1fr);
    gap: 12px;
    padding: 12px 16px;
  }
  .sub-box {
    background: #FFFFFF;
    border: 1px solid #BFDBFE;
    border-radius: 8px;
    padding: 10px 12px;
    text-align: center;
    box-shadow: 0 1px 3px rgba(0,0,0,0.04);
  }
  .sub-box-title {
    font-size: 13.5px;
    font-weight: 800;
    color: #1E3A8A;
    margin-bottom: 4px;
  }
  .sub-box-desc {
    font-size: 11.5px;
    color: #475569;
    line-height: 1.3;
  }
  
  /* 5. Pink */
  .card-5 { border-color: #FBCFE8; background: #FDF2F8; }
  .card-5 .card-header { background: #FCE7F3; }
  .card-5 .badge { background: #DB2777; }
  .card-5 .card-title { color: #9D174D; }
  
  /* 6. Gold/Warm Amber */
  .card-6 { border-color: #FDE68A; background: #FFFBEB; }
  .card-6 .card-header { background: #FEF3C7; }
  .card-6 .badge { background: #D97706; }
  .card-6 .card-title { color: #92400E; }
  
  /* 7. Purple */
  .card-7 { border-color: #DDD6FE; background: #F5F3FF; }
  .card-7 .card-header { background: #EDE9FE; }
  .card-7 .badge { background: #7C3AED; }
  .card-7 .card-title { color: #5B21B6; }
  
  /* 8. Light Blue */
  .card-8 { border-color: #BAE6FD; background: #F0F9FF; }
  .card-8 .card-header { background: #E0F2FE; }
  .card-8 .badge { background: #0284C7; }
  .card-8 .card-title { color: #075985; }
  
  /* 9. Coral / Red */
  .card-9 { border-color: #FECACA; background: #FEF2F2; }
  .card-9 .card-header { background: #FEE2E2; }
  .card-9 .badge { background: #DC2626; }
  .card-9 .card-title { color: #991B1B; }
  
  /* 10. Emerald Teal */
  .card-10 { border-color: #A7F3D0; background: #ECFDF5; }
  .card-10 .card-header { background: #D1FAE5; }
  .card-10 .badge { background: #059669; }
  .card-10 .card-title { color: #065F46; }

  .connector-arrow {
    display: inline-block;
    color: #64748B;
    font-size: 18px;
    font-weight: 900;
  }
</style>
</head>
<body>

  <!-- Top Header Banner -->
  <div class="header-banner">
    <div class="header-title">CIVICCONNECT / JH-INNOVATE</div>
    <div class="header-subtitle">AI & MCP-Powered Civic Grievance Triage, R&D Matchmaking & Resolution Platform</div>
    <div class="header-tags">Multi-Layer AI Forensics &bull; Model Context Protocol (MCP) Decision Tools &bull; Smart Matchmaking &bull; Continuous SLA Audit</div>
  </div>

  <!-- Row 1 (Steps 1, 2, 3) -->
  <div class="row-3-col">
    <!-- Card 1 -->
    <div class="card card-1">
      <div class="card-header">
        <div class="badge">1</div>
        <div class="card-title">Civic Problem Input</div>
      </div>
      <div class="card-body">
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Live Citizen Report:</strong> Web & Mobile Portal</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>GPS Geolocation:</strong> Locked Lat/Long (&plusmn;5m)</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>On-Site Camera Evidence:</strong> Live capture</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Multilingual Text:</strong> Hindi & English inputs</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Zero Login Barrier:</strong> Frictionless access</span></div>
      </div>
    </div>

    <!-- Card 2 -->
    <div class="card card-2">
      <div class="card-header">
        <div class="badge">2</div>
        <div class="card-title">Preprocessing & Validation</div>
      </div>
      <div class="card-body">
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Geofence Validation:</strong> Boundary geocoding</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>EXIF & Resolution Check:</strong> Metadata verification</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Scene Lighting Check:</strong> Underexposure/blur filter</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Text Normalization:</strong> Strips noise & cleans input</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Proximity Clustering:</strong> 200m spatial group</span></div>
      </div>
    </div>

    <!-- Card 3 -->
    <div class="card card-3">
      <div class="card-header">
        <div class="badge">3</div>
        <div class="card-title">Real-Time AI Forensics</div>
      </div>
      <div class="card-body">
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>ELA Tamper Check:</strong> Error Level Analysis</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Deepfake & Screen Check:</strong> Rejects photo of screen</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Noise Variance:</strong> Pixel artifact analysis</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Inconsistency Scoring:</strong> Compression ratio</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Authenticity Score:</strong> 0&ndash;100% integrity index</span></div>
      </div>
    </div>
  </div>

  <!-- Row 2 (Step 4: Wide Span) -->
  <div class="row-wide">
    <div class="card card-4">
      <div class="card-header">
        <div class="badge">4</div>
        <div class="card-title">Multi-Layer AI & MCP Forensic Analysis</div>
      </div>
      <div class="wide-grid">
        <div class="sub-box">
          <div class="sub-box-title">Semantic Duplicate</div>
          <div class="sub-box-desc">Sentence-Transformers cluster recurring neighborhood complaints into Master Problems.</div>
        </div>
        <div class="sub-box">
          <div class="sub-box-title">Domain NLP</div>
          <div class="sub-box-desc">Classifies into 11 civic sectors: Water, Roads, Health, Agri, Sanitation, Energy, etc.</div>
        </div>
        <div class="sub-box">
          <div class="sub-box-title">Urgency & Priority</div>
          <div class="sub-box-desc">Auto-prioritizes hazards: Critical (Safety Risk), High (Disruption), Medium, Low.</div>
        </div>
        <div class="sub-box">
          <div class="sub-box-title">MCP Decision Tools</div>
          <div class="sub-box-desc">Model Context Protocol tools for automated triage, SLA audit & dispatch recommendations.</div>
        </div>
        <div class="sub-box">
          <div class="sub-box-title">Synergy & Feasibility</div>
          <div class="sub-box-desc">Evaluates cross-domain collaboration viability between university and industry proposals.</div>
        </div>
      </div>
    </div>
  </div>

  <!-- Row 3 (Steps 5, 6, 7) -->
  <div class="row-3-col">
    <!-- Card 5 -->
    <div class="card card-5">
      <div class="card-header">
        <div class="badge">5</div>
        <div class="card-title">Partner Matchmaking</div>
      </div>
      <div class="card-body">
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Academic Lab Matching:</strong> Department capability</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Domain Expertise:</strong> Faculty & student teams</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Proximity Ranking:</strong> Prioritizes district colleges</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>CSR Synergy:</strong> Aligns with industry focus</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Compatibility Score:</strong> Multi-factor ranking</span></div>
      </div>
    </div>

    <!-- Card 6 -->
    <div class="card card-6">
      <div class="card-header">
        <div class="badge">6</div>
        <div class="card-title">Proposals & CSR Co-Funding</div>
      </div>
      <div class="card-body">
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>University Proposals:</strong> Engineering R&D plans</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Technical Milestones:</strong> Timeline & budget breakdown</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Industry CSR Grants:</strong> Private co-funding pledges</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>MCP Multi-Proposal:</strong> Automated comparison</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Admin Gateways:</strong> Gate 1 & 2 review & approval</span></div>
      </div>
    </div>

    <!-- Card 7 -->
    <div class="card card-7">
      <div class="card-header">
        <div class="badge">7</div>
        <div class="card-title">Lifecycle & Work Order Engine</div>
      </div>
      <div class="card-body">
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Spring Boot 3.4:</strong> Robust REST API & Gateway</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>State Machine:</strong> NEW &rarr; ROUTED &rarr; COMPLETED</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Role-Based Access:</strong> Stateless JWT security</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Task Allocation:</strong> Real-time work assignments</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Milestone Ledger:</strong> Document & progress store</span></div>
      </div>
    </div>
  </div>

  <!-- Row 4 (Steps 8, 9, 10) -->
  <div class="row-3-col">
    <!-- Card 8 -->
    <div class="card card-8">
      <div class="card-header">
        <div class="badge">8</div>
        <div class="card-title">Dynamic SLA Delay Radar</div>
      </div>
      <div class="card-body">
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Milestone Velocity:</strong> Real-time progress tracker</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Delay Risk Score (0&ndash;100):</strong> Predictive index</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>SLA Breach Forecast:</strong> Early risk detection</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Trajectory Health:</strong> Green / Yellow / Red status</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Explainable Alerts:</strong> Root-cause risk factors</span></div>
      </div>
    </div>

    <!-- Card 9 -->
    <div class="card card-9">
      <div class="card-header">
        <div class="badge">9</div>
        <div class="card-title">AI & Admin Decision Engine</div>
      </div>
      <div class="card-body">
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Decision Gateways:</strong> Gate 1 (Triage) & Gate 2 (Collab)</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>APPROVE (0&ndash;30):</strong> Low risk, direct execution</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>REVIEW (31&ndash;70):</strong> Request info / revision</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>ESCALATE (71&ndash;100):</strong> Warning & admin directive</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Escrow Release:</strong> Authorizes CSR fund tranches</span></div>
      </div>
    </div>

    <!-- Card 10 -->
    <div class="card card-10">
      <div class="card-header">
        <div class="badge">10</div>
        <div class="card-title">Resolution & Field Audit</div>
      </div>
      <div class="card-body">
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Verified Resolution:</strong> Completed physical works</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Photo Evidence Audit:</strong> Before/After verification</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Citizen Redressal:</strong> Live tracking update</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>MongoDB Record:</strong> Permanent immutable impact log</span></div>
        <div class="bullet-item"><span class="bullet-dot">&bull;</span><span><strong>Public Dashboards:</strong> Open state governance SLA</span></div>
      </div>
    </div>
  </div>

</body>
</html>
"""
    with open(output_html_path, 'w', encoding='utf-8') as f:
        f.write(html_content)
    print(f"HTML created at {output_html_path}")

def create_pptx_version(output_pptx_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)
    
    # Background color
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(241, 245, 249)
    bg.line.color.rgb = RGBColor(241, 245, 249)
    
    # Header Banner
    header = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.3), Inches(0.2), Inches(12.733), Inches(0.95))
    header.fill.solid()
    header.fill.fore_color.rgb = RGBColor(10, 25, 47)
    header.line.color.rgb = RGBColor(30, 58, 138)
    
    tf = header.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.1)
    tf.margin_top = Inches(0.08)
    tf.margin_right = Inches(0.1)
    tf.margin_bottom = Inches(0.08)
    
    p1 = tf.paragraphs[0]
    p1.text = "CIVICCONNECT / JH-INNOVATE"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.bold = True
    p1.font.size = Pt(17)
    p1.font.color.rgb = RGBColor(255, 255, 255)
    
    p2 = tf.add_paragraph()
    p2.text = "AI & MCP-Powered Civic Grievance Triage, R&D Matchmaking & Resolution Platform"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.bold = True
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = RGBColor(147, 197, 253)
    
    p3 = tf.add_paragraph()
    p3.text = "Multi-Layer AI Forensics • Model Context Protocol (MCP) Decision Tools • Smart Matchmaking • Continuous SLA Audit"
    p3.alignment = PP_ALIGN.CENTER
    p3.font.size = Pt(8.5)
    p3.font.color.rgb = RGBColor(226, 232, 240)
    
    # Helper to create a 3-column card
    def add_card(x, y, w, h, num, title, bullets, border_rgb, bg_rgb, badge_rgb, title_rgb):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_rgb
        card.line.color.rgb = border_rgb
        card.line.width = Pt(1.5)
        
        # Badge
        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.12), y + Inches(0.08), Inches(0.32), Inches(0.32))
        badge.fill.solid()
        badge.fill.fore_color.rgb = badge_rgb
        badge.line.color.rgb = badge_rgb
        btf = badge.text_frame
        btf.margin_left = btf.margin_top = btf.margin_right = btf.margin_bottom = Inches(0)
        bp = btf.paragraphs[0]
        bp.text = str(num)
        bp.alignment = PP_ALIGN.CENTER
        bp.font.bold = True
        bp.font.size = Pt(11)
        bp.font.color.rgb = RGBColor(255, 255, 255)
        
        # Text Frame
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.48)
        ctf.margin_top = Inches(0.08)
        ctf.margin_right = Inches(0.1)
        ctf.margin_bottom = Inches(0.05)
        
        tp = ctf.paragraphs[0]
        tp.text = title.upper()
        tp.font.bold = True
        tp.font.size = Pt(10)
        tp.font.color.rgb = title_rgb
        
        for b in bullets:
            bp = ctf.add_paragraph()
            bp.text = f"• {b}"
            bp.font.size = Pt(8.5)
            bp.font.color.rgb = RGBColor(30, 41, 59)
            bp.space_before = Pt(2)
            
    col_w = Inches(4.08)
    gap = Inches(0.24)
    x1 = Inches(0.3)
    x2 = x1 + col_w + gap
    x3 = x2 + col_w + gap
    
    # ROW 1 (y = 1.25, h = 1.35)
    r1_y = Inches(1.25)
    r1_h = Inches(1.35)
    add_card(x1, r1_y, col_w, r1_h, 1, "Civic Problem Input", [
        "Live Citizen Report: Web & Mobile Portal",
        "GPS Geolocation Lock: (±5m precision)",
        "On-Site Live Camera Evidence",
        "Multilingual Inputs (Hindi / English)",
        "Zero Mandatory Login Barrier"
    ], RGBColor(252, 211, 77), RGBColor(255, 251, 235), RGBColor(217, 119, 6), RGBColor(146, 64, 14))

    add_card(x2, r1_y, col_w, r1_h, 2, "Preprocessing & Validation", [
        "Geofence & Boundary Verification",
        "EXIF & Image Quality Check",
        "Scene Lighting & Blur Filtering",
        "Multilingual Text Normalization",
        "Proximity Radius Clustering (200m)"
    ], RGBColor(167, 243, 208), RGBColor(236, 253, 245), RGBColor(5, 150, 105), RGBColor(6, 95, 70))

    add_card(x3, r1_y, col_w, r1_h, 3, "Real-Time AI Forensics", [
        "ELA Tamper Analysis (Pixel Residuals)",
        "Deepfake & Screen Photo Detection",
        "Noise Variance & Edge Analysis",
        "Compression Ratio Scoring",
        "Authenticity Score (0-100%)"
    ], RGBColor(253, 186, 116), RGBColor(255, 247, 237), RGBColor(234, 88, 12), RGBColor(154, 52, 18))

    # ROW 2 (Wide Step 4: y = 2.7, h = 1.25)
    r2_y = Inches(2.7)
    r2_h = Inches(1.25)
    w_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.3), r2_y, Inches(12.733), r2_h)
    w_card.fill.solid()
    w_card.fill.fore_color.rgb = RGBColor(239, 246, 255)
    w_card.line.color.rgb = RGBColor(147, 197, 253)
    w_card.line.width = Pt(1.5)
    
    # Badge 4
    b4 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.42), r2_y + Inches(0.08), Inches(0.32), Inches(0.32))
    b4.fill.solid()
    b4.fill.fore_color.rgb = RGBColor(37, 99, 235)
    b4.line.color.rgb = RGBColor(37, 99, 235)
    b4_tf = b4.text_frame
    b4_tf.margin_left = b4_tf.margin_top = b4_tf.margin_right = b4_tf.margin_bottom = Inches(0)
    b4_p = b4_tf.paragraphs[0]
    b4_p.text = "4"
    b4_p.alignment = PP_ALIGN.CENTER
    b4_p.font.bold = True
    b4_p.font.size = Pt(11)
    b4_p.font.color.rgb = RGBColor(255, 255, 255)
    
    # Title 4
    t4 = slide.shapes.add_textbox(Inches(0.8), r2_y + Inches(0.05), Inches(11.5), Inches(0.3))
    t4_tf = t4.text_frame
    t4_tf.margin_left = t4_tf.margin_top = t4_tf.margin_right = t4_tf.margin_bottom = Inches(0)
    t4_p = t4_tf.paragraphs[0]
    t4_p.text = "MULTI-LAYER AI & MCP FORENSIC ANALYSIS"
    t4_p.font.bold = True
    t4_p.font.size = Pt(10.5)
    t4_p.font.color.rgb = RGBColor(30, 64, 175)
    
    # 5 Sub-boxes
    sub_w = Inches(2.4)
    sub_gap = Inches(0.12)
    sub_y = r2_y + Inches(0.38)
    sub_h = Inches(0.78)
    
    sub_items = [
        ("Semantic Duplicate", "Sentence-Transformers group neighborhood reports into Master Problem."),
        ("Domain NLP", "Classifies into 11 civic sectors: Water, Roads, Health, Agri, Sanitation..."),
        ("Urgency Scoring", "Auto-prioritizes: Critical (Hazard), High (Disruption), Medium, Low."),
        ("MCP Decision Tools", "Model Context Protocol tools for automated triage, SLA audit & dispatch."),
        ("Synergy Check", "Evaluates feasibility & cross-domain alignment between Univ & CSR.")
    ]
    
    for i, (stitle, sdesc) in enumerate(sub_items):
        sx = Inches(0.42) + i * (sub_w + sub_gap)
        sbox = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx, sub_y, sub_w, sub_h)
        sbox.fill.solid()
        sbox.fill.fore_color.rgb = RGBColor(255, 255, 255)
        sbox.line.color.rgb = RGBColor(191, 219, 254)
        stf = sbox.text_frame
        stf.word_wrap = True
        stf.margin_left = stf.margin_right = Inches(0.08)
        stf.margin_top = stf.margin_bottom = Inches(0.05)
        
        sp1 = stf.paragraphs[0]
        sp1.text = stitle
        sp1.font.bold = True
        sp1.font.size = Pt(8.5)
        sp1.font.color.rgb = RGBColor(30, 58, 138)
        
        sp2 = stf.add_paragraph()
        sp2.text = sdesc
        sp2.font.size = Pt(7.2)
        sp2.font.color.rgb = RGBColor(71, 85, 105)
        sp2.space_before = Pt(2)
        
    # ROW 3 (Steps 5, 6, 7: y = 4.05, h = 1.45)
    r3_y = Inches(4.05)
    r3_h = Inches(1.45)
    add_card(x1, r3_y, col_w, r3_h, 5, "Partner Matchmaking", [
        "Academic Lab Capability Matching",
        "University Department Expertise",
        "District Proximity Ranking",
        "Industry CSR Domain Alignment",
        "Match Compatibility Score Generation"
    ], RGBColor(251, 207, 232), RGBColor(253, 242, 248), RGBColor(219, 39, 119), RGBColor(157, 23, 77))

    add_card(x2, r3_y, col_w, r3_h, 6, "Proposals & CSR Co-Funding", [
        "University R&D Engineering Proposals",
        "Milestone Roadmaps & Technical Plan",
        "Industry CSR Co-Funding Pledges",
        "MCP Multi-Proposal Synergy Compare",
        "Admin Gate 1 & 2 Review & Approval"
    ], RGBColor(253, 230, 138), RGBColor(255, 251, 235), RGBColor(217, 119, 6), RGBColor(146, 64, 14))

    add_card(x3, r3_y, col_w, r3_h, 7, "Lifecycle & Work Order Engine", [
        "Spring Boot 3.4 REST API & Gateway",
        "State Machine (NEW to COMPLETED)",
        "Stateless JWT Security & RBAC",
        "Real-Time Task & Team Allocation",
        "Document & Milestone Ledger"
    ], RGBColor(221, 214, 254), RGBColor(245, 243, 255), RGBColor(124, 58, 237), RGBColor(91, 33, 182))

    # ROW 4 (Steps 8, 9, 10: y = 5.6, h = 1.45)
    r4_y = Inches(5.6)
    r4_h = Inches(1.45)
    add_card(x1, r4_y, col_w, r4_h, 8, "Dynamic SLA Delay Radar", [
        "Milestone Velocity Real-Time Tracker",
        "Delay Risk Score (0-100)",
        "SLA Breach Probability Forecast",
        "Trajectory Health (Green / Amber / Red)",
        "Explainable Risk Factor Directives"
    ], RGBColor(186, 230, 253), RGBColor(240, 249, 255), RGBColor(2, 132, 199), RGBColor(7, 89, 133))

    add_card(x2, r4_y, col_w, r4_h, 9, "AI & Admin Decision Engine", [
        "Admin Decision Gateways (1 & 2)",
        "APPROVE (0-30): Low Risk Execution",
        "REVIEW (31-70): Info Request / Refine",
        "ESCALATE (71-100): Delay Warning",
        "CSR Fund Release & Joint Work Orders"
    ], RGBColor(254, 202, 202), RGBColor(254, 242, 242), RGBColor(220, 38, 38), RGBColor(153, 27, 27))

    add_card(x3, r4_y, col_w, r4_h, 10, "Resolution & Field Audit", [
        "Verified On-Site Field Resolution",
        "Before / After Photo Evidence Audit",
        "Citizen Grievance Redressal Update",
        "Permanent Immutable MongoDB Log",
        "Open Governance SLA Dashboards"
    ], RGBColor(167, 243, 208), RGBColor(236, 253, 245), RGBColor(5, 150, 105), RGBColor(6, 95, 70))

    prs.save(output_pptx_path)
    print(f"PowerPoint Presentation saved at {output_pptx_path}")

if __name__ == '__main__':
    root_dir = r"c:\Users\kesav\OneDrive\Desktop\CivicConnect"
    html_path = os.path.join(root_dir, "CIVICCONNECT_Architecture.html")
    pptx_path = os.path.join(root_dir, "CIVICCONNECT_Architecture.pptx")
    img_path = os.path.join(root_dir, "CIVICCONNECT_Architecture_Human.png")
    dest_pub = os.path.join(root_dir, "client", "public", "CIVICCONNECT_Architecture_Human.png")
    
    create_html_version(html_path)
    create_pptx_version(pptx_path)
