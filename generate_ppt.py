import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]
    
    # ----------------------------------------------------
    # SLIDE 1: TECHNICAL APPROACH (1:1 SIH Template)
    # ----------------------------------------------------
    slide = prs.slides.add_slide(blank_slide_layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    
    # Header
    logo_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(3.0), Inches(0.6))
    tf = logo_box.text_frame
    p = tf.paragraphs[0]
    p.text = "CIVIC CONNECT"
    p.font.name = "Arial Black"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = RGBColor(24, 43, 73)
    p2 = tf.add_paragraph()
    p2.text = "Govt. of Jharkhand Civic Innovation"
    p2.font.name = "Arial"
    p2.font.size = Pt(9)
    p2.font.color.rgb = RGBColor(100, 116, 139)
    
    title_box = slide.shapes.add_textbox(Inches(3.8), Inches(0.18), Inches(5.8), Inches(0.7))
    p_title = title_box.text_frame.paragraphs[0]
    p_title.text = "TECHNICAL APPROACH"
    p_title.alignment = PP_ALIGN.CENTER
    p_title.font.name = "Arial Black"
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(15, 23, 42)
    
    sih_box = slide.shapes.add_textbox(Inches(10.2), Inches(0.15), Inches(2.7), Inches(0.7))
    p_sih = sih_box.text_frame.paragraphs[0]
    p_sih.text = "SMART INDIA HACKATHON 2025"
    p_sih.alignment = PP_ALIGN.RIGHT
    p_sih.font.name = "Arial"
    p_sih.font.size = Pt(11)
    p_sih.font.bold = True
    p_sih.font.color.rgb = RGBColor(225, 112, 85)
    
    # Left Column (Pure Software Stack)
    left_x = Inches(0.4)
    left_w = Inches(2.7)
    
    # 1. System Environment & Runtime
    hw_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, Inches(0.95), left_w, Inches(1.15))
    hw_card.fill.solid()
    hw_card.fill.fore_color.rgb = RGBColor(248, 250, 252)
    hw_card.line.color.rgb = RGBColor(203, 213, 225)
    
    pill1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x + Inches(0.1), Inches(0.85), Inches(2.5), Inches(0.28))
    pill1.fill.solid()
    pill1.fill.fore_color.rgb = RGBColor(59, 130, 246)
    pill1.line.fill.background()
    p = pill1.text_frame.paragraphs[0]
    p.text = "System Environment & Runtime"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    tf = hw_card.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.22)
    tf.margin_left = Inches(0.1)
    p = tf.paragraphs[0]
    p.text = "• Runtime: Java 17 (JDK) & Python 3.10+\n• Client: Web & Mobile Responsive (React 18)\n• API Gateway: Spring Boot 3.4 (Port 5000)\n• AI Engine: FastAPI ASGI Server (Port 8000)"
    p.font.size = Pt(7.8)
    p.font.color.rgb = RGBColor(30, 41, 59)
    
    # 2. Frameworks & Libraries
    sw_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, Inches(2.25), left_w, Inches(1.2))
    sw_card.fill.solid()
    sw_card.fill.fore_color.rgb = RGBColor(248, 250, 252)
    sw_card.line.color.rgb = RGBColor(203, 213, 225)
    
    pill2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x + Inches(0.1), Inches(2.15), Inches(2.5), Inches(0.28))
    pill2.fill.solid()
    pill2.fill.fore_color.rgb = RGBColor(79, 70, 229)
    pill2.line.fill.background()
    p = pill2.text_frame.paragraphs[0]
    p.text = "Frameworks & Libraries"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    tf = sw_card.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.22)
    tf.margin_left = Inches(0.1)
    p = tf.paragraphs[0]
    p.text = "• Java 17 + Spring Boot 3.4.3 (REST Gateway)\n• Python 3.10 + FastAPI (AI Microservice)\n• HuggingFace XLM-RoBERTa (Indic NLP)\n• PyTorch / OpenCV / PIL (ELA Forensics)\n• React 18 + Vite + Tailwind / Leaflet GIS"
    p.font.size = Pt(7.8)
    p.font.color.rgb = RGBColor(30, 41, 59)
    
    # 3. Storage Technologies
    st_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, Inches(3.6), left_w, Inches(1.05))
    st_card.fill.solid()
    st_card.fill.fore_color.rgb = RGBColor(248, 250, 252)
    st_card.line.color.rgb = RGBColor(203, 213, 225)
    
    pill3 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x + Inches(0.1), Inches(3.5), Inches(2.5), Inches(0.28))
    pill3.fill.solid()
    pill3.fill.fore_color.rgb = RGBColor(14, 165, 233)
    pill3.line.fill.background()
    p = pill3.text_frame.paragraphs[0]
    p.text = "Storage Technologies"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    tf = st_card.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.22)
    tf.margin_left = Inches(0.1)
    p = tf.paragraphs[0]
    p.text = "• MongoDB 7.0 (Problems, Solutions, Users)\n• Cloudinary / GridFS (Media Evidence Photos)\n• Vector Embeddings (MiniLM-L6 Cosine Cache)\n• IndexedDB (Offline Submission Sync)"
    p.font.size = Pt(7.8)
    p.font.color.rgb = RGBColor(30, 41, 59)
    
    # 4. Integration & Protocols
    cl_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, Inches(4.8), left_w, Inches(1.05))
    cl_card.fill.solid()
    cl_card.fill.fore_color.rgb = RGBColor(248, 250, 252)
    cl_card.line.color.rgb = RGBColor(203, 213, 225)
    
    pill4 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x + Inches(0.1), Inches(4.7), Inches(2.5), Inches(0.28))
    pill4.fill.solid()
    pill4.fill.fore_color.rgb = RGBColor(2, 132, 199)
    pill4.line.fill.background()
    p = pill4.text_frame.paragraphs[0]
    p.text = "Integration & Protocols"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    tf = cl_card.text_frame
    tf.word_wrap = True
    tf.margin_top = Inches(0.22)
    tf.margin_left = Inches(0.1)
    p = tf.paragraphs[0]
    p.text = "• Model Context Protocol (72 Production Tools)\n• RESTful JSON HTTP APIs & WebSockets\n• JWT Token Auth & BCrypt Password Hash\n• GitHub Actions Automated CI/CD Pipelines"
    p.font.size = Pt(7.8)
    p.font.color.rgb = RGBColor(30, 41, 59)
    
    # Bottom Left Yellow Box
    yellow_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, Inches(5.95), Inches(4.4), Inches(1.15))
    yellow_card.fill.solid()
    yellow_card.fill.fore_color.rgb = RGBColor(254, 240, 138)
    yellow_card.line.color.rgb = RGBColor(234, 179, 8)
    yellow_card.line.width = Pt(1.5)
    
    tf_y = yellow_card.text_frame
    tf_y.word_wrap = True
    tf_y.margin_left = Inches(0.12)
    tf_y.margin_top = Inches(0.06)
    
    items = [
        "• Scope: Ingest civic grievances, match university R&D with CSR grants, enforce SLAs",
        "• Design: 3-tier microservice architecture (React 18 + Spring Boot 3.4 + FastAPI AI + MCP)",
        "• Development: Multilingual Indic NLP, ELA photo authenticity & MCDA matchmaking models",
        "• Deployment: Containerized microservices, MongoDB persistence & JWT role-based security",
        "• Maintenance: Real-time SLA tracking, multi-district analytics & tamper-evident audit logs"
    ]
    for i, item in enumerate(items):
        p = tf_y.paragraphs[0] if i == 0 else tf_y.add_paragraph()
        p.text = item
        p.font.size = Pt(7.4)
        p.font.color.rgb = RGBColor(113, 63, 18)
        if i == 0:
            p.font.bold = True

    # Center Pipeline Boxes
    # 1. Ingestion
    col1_x = Inches(3.25)
    c1_w = Inches(1.7)
    inp_group = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col1_x, Inches(0.95), c1_w, Inches(4.85))
    inp_group.fill.solid()
    inp_group.fill.fore_color.rgb = RGBColor(241, 245, 249)
    inp_group.line.color.rgb = RGBColor(203, 213, 225)
    
    hdr1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col1_x + Inches(0.08), Inches(1.02), c1_w - Inches(0.16), Inches(0.32))
    hdr1.fill.solid()
    hdr1.fill.fore_color.rgb = RGBColor(30, 41, 59)
    hdr1.line.fill.background()
    p = hdr1.text_frame.paragraphs[0]
    p.text = "1. Data Ingestion"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    sub_inputs = [
        ("📱 Citizen Portal", "Bilingual text (Hindi/Eng)\nGeo-tagged image upload"),
        ("🖥️ Admin Command", "Live triage, heatmaps &\nDepartment routing"),
        ("🎓 University Lab", "Technical R&D proposals,\nBOM budget & timeline"),
        ("🏭 Industry CSR Hub", "CSR grant allocation &\nMilestone co-funding")
    ]
    for i, (stitle, sdesc) in enumerate(sub_inputs):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col1_x + Inches(0.08), Inches(1.42 + i*1.05), c1_w - Inches(0.16), Inches(0.95))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(255, 255, 255)
        box.line.color.rgb = RGBColor(226, 232, 240)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.06)
        tf.margin_left = Inches(0.06)
        p = tf.paragraphs[0]
        p.text = stitle
        p.font.size = Pt(8.2)
        p.font.bold = True
        p.font.color.rgb = RGBColor(15, 23, 42)
        p2 = tf.add_paragraph()
        p2.text = sdesc
        p2.font.size = Pt(7.2)
        p2.font.color.rgb = RGBColor(100, 116, 139)

    # 2. Cleansing
    col2_x = Inches(5.1)
    c2_w = Inches(1.8)
    cln_group = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col2_x, Inches(0.95), c2_w, Inches(4.85))
    cln_group.fill.solid()
    cln_group.fill.fore_color.rgb = RGBColor(240, 253, 250)
    cln_group.line.color.rgb = RGBColor(153, 246, 228)
    
    hdr2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col2_x + Inches(0.08), Inches(1.02), c2_w - Inches(0.16), Inches(0.32))
    hdr2.fill.solid()
    hdr2.fill.fore_color.rgb = RGBColor(13, 148, 136)
    hdr2.line.fill.background()
    p = hdr2.text_frame.paragraphs[0]
    p.text = "2. Cleansing & Pre-filter"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    sub_clean = [
        ("🔍 ELA Image Forensics", "Error Level Analysis &\nAI fake / tamper check"),
        ("🌐 Indic Multilingual", "Hindi/English tokenization,\nregional dialect normalize"),
        ("📑 Semantic Dedup", "MiniLM cosine matching\nmerges duplicate tickets"),
        ("🛡️ Input Sanitizer", "Regex security filter &\nNoSQL injection guard")
    ]
    for i, (stitle, sdesc) in enumerate(sub_clean):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col2_x + Inches(0.08), Inches(1.42 + i*1.05), c2_w - Inches(0.16), Inches(0.95))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(255, 255, 255)
        box.line.color.rgb = RGBColor(204, 251, 241)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.06)
        tf.margin_left = Inches(0.06)
        p = tf.paragraphs[0]
        p.text = stitle
        p.font.size = Pt(8.2)
        p.font.bold = True
        p.font.color.rgb = RGBColor(17, 94, 89)
        p2 = tf.add_paragraph()
        p2.text = sdesc
        p2.font.size = Pt(7.2)
        p2.font.color.rgb = RGBColor(71, 85, 105)

    # 3. AI / ML
    col3_x = Inches(7.05)
    c3_w = Inches(2.2)
    ai_group = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col3_x, Inches(0.95), c3_w, Inches(4.85))
    ai_group.fill.solid()
    ai_group.fill.fore_color.rgb = RGBColor(238, 242, 255)
    ai_group.line.color.rgb = RGBColor(199, 210, 254)
    ai_group.line.width = Pt(1.5)
    
    hdr3 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col3_x + Inches(0.08), Inches(1.02), c3_w - Inches(0.16), Inches(0.32))
    hdr3.fill.solid()
    hdr3.fill.fore_color.rgb = RGBColor(67, 56, 202)
    hdr3.line.fill.background()
    p = hdr3.text_frame.paragraphs[0]
    p.text = "3. AI / ML & MCP Engine"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    sub_ai = [
        ("🤖 XLM-RoBERTa NLP", "Auto domain classifier &\nSLA urgency computation"),
        ("👁️ Vision CNN Analyzer", "Multi-label scene check for\npotholes, sewage & garbage"),
        ("🤝 Institutional Matchmaker", "MCDA algorithm matching\nUniversity labs with CSR"),
        ("⚡ 72 MCP Production Tools", "Assisted decision support for\nGate 1, Gate 2 & Gate 3")
    ]
    for i, (stitle, sdesc) in enumerate(sub_ai):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col3_x + Inches(0.08), Inches(1.42 + i*1.05), c3_w - Inches(0.16), Inches(0.95))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(255, 255, 255)
        box.line.color.rgb = RGBColor(224, 231, 255)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.06)
        tf.margin_left = Inches(0.06)
        p = tf.paragraphs[0]
        p.text = stitle
        p.font.size = Pt(8.2)
        p.font.bold = True
        p.font.color.rgb = RGBColor(49, 46, 129)
        p2 = tf.add_paragraph()
        p2.text = sdesc
        p2.font.size = Pt(7.2)
        p2.font.color.rgb = RGBColor(71, 85, 105)

    # 4. Portals & Backend
    col4_x = Inches(9.4)
    c4_w = Inches(1.8)
    be_group = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col4_x, Inches(1.95), c4_w, Inches(3.85))
    be_group.fill.solid()
    be_group.fill.fore_color.rgb = RGBColor(248, 250, 252)
    be_group.line.color.rgb = RGBColor(203, 213, 225)
    
    hdr4 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col4_x + Inches(0.08), Inches(2.02), c4_w - Inches(0.16), Inches(0.32))
    hdr4.fill.solid()
    hdr4.fill.fore_color.rgb = RGBColor(30, 41, 59)
    hdr4.line.fill.background()
    p = hdr4.text_frame.paragraphs[0]
    p.text = "4. Backend & Portals"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    sub_portals = [
        ("🖥️ Admin Command", "Triaging, GIS heatmaps, SLA\nescalation & work orders"),
        ("🎓 University Proposals", "Proposal drafting, BOM &\nPrototype milestone tracker"),
        ("🏭 Industry CSR Portal", "CSR co-funding allocation &\nMilestone escrow release")
    ]
    for i, (stitle, sdesc) in enumerate(sub_portals):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, col4_x + Inches(0.08), Inches(2.42 + i*1.1), c4_w - Inches(0.16), Inches(1.0))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(255, 255, 255)
        box.line.color.rgb = RGBColor(226, 232, 240)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.06)
        tf.margin_left = Inches(0.06)
        p = tf.paragraphs[0]
        p.text = stitle
        p.font.size = Pt(8.2)
        p.font.bold = True
        p.font.color.rgb = RGBColor(15, 23, 42)
        p2 = tf.add_paragraph()
        p2.text = sdesc
        p2.font.size = Pt(7.2)
        p2.font.color.rgb = RGBColor(100, 116, 139)

    # Top Right Product Status Post-it
    status_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.35), Inches(0.95), Inches(1.65), Inches(1.8))
    status_card.fill.solid()
    status_card.fill.fore_color.rgb = RGBColor(254, 243, 199)
    status_card.line.color.rgb = RGBColor(245, 158, 11)
    status_card.line.width = Pt(1.5)
    
    pin = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(12.75), Inches(0.85), Inches(0.22), Inches(0.22))
    pin.fill.solid()
    pin.fill.fore_color.rgb = RGBColor(220, 38, 38)
    pin.line.fill.background()
    
    tf_s = status_card.text_frame
    tf_s.word_wrap = True
    tf_s.margin_top = Inches(0.1)
    p = tf_s.paragraphs[0]
    p.text = "Product Status"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(146, 64, 14)
    p2 = tf_s.add_paragraph()
    p2.text = "100% Core System Built & Verified!\n(54/54 Test Suite Passed)\nProduction Ready"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(8.5)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(30, 64, 175)

    # Security Box
    sec_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.35), Inches(2.9), Inches(1.65), Inches(1.9))
    sec_box.fill.solid()
    sec_box.fill.fore_color.rgb = RGBColor(245, 243, 255)
    sec_box.line.color.rgb = RGBColor(221, 214, 254)
    tf_sec = sec_box.text_frame
    tf_sec.word_wrap = True
    p = tf_sec.paragraphs[0]
    p.text = "🔒 Security & Auth"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(109, 40, 217)
    p2 = tf_sec.add_paragraph()
    p2.text = "• JWT + BCrypt Auth\n• Role-Based Guard (RBAC)\n• AES-256 Data Encryption\n• SHA-256 Audit Trail\n• Human-In-The-Loop AI"
    p2.font.size = Pt(7.2)
    p2.font.color.rgb = RGBColor(68, 64, 60)

    # Readiness Badges
    badge_data = [
        ("7", "TECHNOLOGY\nREADINESS LEVEL", RGBColor(14, 116, 144), Inches(4.95)),
        ("6", "DEPLOYMENT\nREADINESS LEVEL", RGBColor(161, 98, 7), Inches(5.65)),
        ("5", "INVESTMENT\nREADINESS LEVEL", RGBColor(30, 58, 138), Inches(6.35))
    ]
    for num, label, col, y_pos in badge_data:
        bx = Inches(11.35)
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, bx, y_pos, Inches(0.52), Inches(0.52))
        circ.fill.solid()
        circ.fill.fore_color.rgb = col
        circ.line.fill.background()
        p = circ.text_frame.paragraphs[0]
        p.text = num
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        
        lbl = slide.shapes.add_textbox(bx + Inches(0.56), y_pos - Inches(0.05), Inches(1.3), Inches(0.6))
        tf = lbl.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = label
        p.font.size = Pt(6.8)
        p.font.bold = True
        p.font.color.rgb = col

    # Footer Slide 1
    footer = slide.shapes.add_textbox(Inches(4.9), Inches(7.15), Inches(3.6), Inches(0.3))
    p = footer.text_frame.paragraphs[0]
    p.text = "@SIH Idea submission - CIVIC CONNECT"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = RGBColor(100, 116, 139)


    # ----------------------------------------------------
    # SLIDE 2: TECHNOLOGY STACK (Matching User Reference)
    # ----------------------------------------------------
    slide_tech = prs.slides.add_slide(blank_slide_layout)
    slide_tech.background.fill.solid()
    slide_tech.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    
    # Outer decorative container
    border_box = slide_tech.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(0.3), Inches(12.533), Inches(6.9))
    border_box.fill.background()
    border_box.line.color.rgb = RGBColor(203, 213, 225)
    border_box.line.width = Pt(1.5)
    
    # Main Header
    t_box = slide_tech.shapes.add_textbox(Inches(3.5), Inches(0.45), Inches(6.333), Inches(0.7))
    p = t_box.text_frame.paragraphs[0]
    p.text = "Technology Stack"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Georgia"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.underline = True
    p.font.color.rgb = RGBColor(24, 43, 73)
    
    # 4 Rows of Tech Stack
    tech_sections = [
        ("Frontend & UI", [
            ("⚛️ React 18", RGBColor(6, 182, 212)),
            ("⚡ Vite", RGBColor(168, 85, 247)),
            ("🟨 JavaScript", RGBColor(234, 179, 8)),
            ("🌊 Tailwind CSS", RGBColor(14, 165, 233)),
            ("🍃 Leaflet GIS", RGBColor(34, 197, 94)),
            ("🅱️ Bootstrap 5", RGBColor(147, 51, 234))
        ], Inches(1.3)),
        
        ("Backend & Gateway", [
            ("☕ Java 17", RGBColor(234, 88, 12)),
            ("🍃 Spring Boot 3.4", RGBColor(22, 163, 74)),
            ("🔐 JWT Auth", RGBColor(219, 39, 119)),
            ("🪶 Maven", RGBColor(185, 28, 28)),
            ("🌐 RESTful API", RGBColor(37, 99, 235))
        ], Inches(2.7)),
        
        ("AI / ML & Intelligence", [
            ("🐍 Python 3.10", RGBColor(30, 64, 175)),
            ("⚡ FastAPI", RGBColor(13, 148, 136)),
            ("🔥 PyTorch", RGBColor(234, 88, 12)),
            ("🤗 Hugging Face (RoBERTa)", RGBColor(217, 119, 6)),
            ("👁️ OpenCV (ELA Forensics)", RGBColor(220, 38, 38)),
            ("🤖 MCP Server (72 Tools)", RGBColor(79, 70, 229))
        ], Inches(4.1)),
        
        ("Database & DevOps", [
            ("🍃 MongoDB 7.0", RGBColor(22, 163, 74)),
            ("☁️ Cloudinary", RGBColor(2, 132, 199)),
            ("🐳 Docker", RGBColor(3, 105, 161)),
            ("🐙 GitHub Actions", RGBColor(15, 23, 42)),
            ("📊 MiniLM-L6 Vector Cache", RGBColor(124, 58, 237))
        ], Inches(5.5))
    ]
    
    for label, items, y_pos in tech_sections:
        # Category label on left
        lbl_box = slide_tech.shapes.add_textbox(Inches(0.6), y_pos, Inches(2.2), Inches(1.0))
        p = lbl_box.text_frame.paragraphs[0]
        p.text = label
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = RGBColor(15, 23, 42)
        
        # Badges/cards horizontally
        x_start = Inches(2.9)
        spacing = Inches(1.55)
        for i, (name, col) in enumerate(items):
            bx = x_start + i * spacing
            card = slide_tech.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, y_pos, Inches(1.42), Inches(0.95))
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(248, 250, 252)
            card.line.color.rgb = RGBColor(226, 232, 240)
            card.line.width = Pt(1)
            
            tf_b = card.text_frame
            tf_b.word_wrap = True
            tf_b.margin_top = Inches(0.12)
            tf_b.margin_left = Inches(0.05)
            tf_b.margin_right = Inches(0.05)
            
            p = tf_b.paragraphs[0]
            p.text = name
            p.alignment = PP_ALIGN.CENTER
            p.font.name = "Arial"
            p.font.size = Pt(9.5)
            p.font.bold = True
            p.font.color.rgb = col
            
    # Save presentation
    output_path = "CivicConnect_Technical_Approach_Architecture.pptx"
    prs.save(output_path)
    print(f"Updated presentation saved to: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_presentation()
