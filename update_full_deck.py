import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def update_full_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: TECHNICAL APPROACH (SIH 1:1 Template)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    slide1.background.fill.solid()
    slide1.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    # Header
    brand = slide1.shapes.add_textbox(Inches(0.4), Inches(0.18), Inches(3.2), Inches(0.65))
    p = brand.text_frame.paragraphs[0]
    p.text = "CIVIC CONNECT"
    p.font.name = "Arial Black"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 44, 89)
    p2 = brand.text_frame.add_paragraph()
    p2.text = "Govt. of Jharkhand Civic Innovation"
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = RGBColor(100, 116, 139)

    title1 = slide1.shapes.add_textbox(Inches(3.8), Inches(0.15), Inches(5.8), Inches(0.7))
    p = title1.text_frame.paragraphs[0]
    p.text = "TECHNICAL APPROACH"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial Black"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 23, 42)

    sih1 = slide1.shapes.add_textbox(Inches(10.2), Inches(0.15), Inches(2.7), Inches(0.7))
    p = sih1.text_frame.paragraphs[0]
    p.text = "SMART INDIA HACKATHON 2025"
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(225, 112, 85)

    # Left Column
    left_x = Inches(0.4)
    left_w = Inches(2.7)
    specs = [
        ("System Environment", "• Java 17 & Python 3.10+\n• React 18 & Vite Client\n• Spring Boot 3.4 (Port 5000)\n• FastAPI AI (Port 8000)", Inches(0.95), Inches(1.15), RGBColor(59, 130, 246)),
        ("Frameworks & Libraries", "• Spring Boot 3.4 & Security\n• FastAPI & PyTorch\n• Sentence-Transformers\n• Pillow (ELA Forensics)", Inches(2.25), Inches(1.15), RGBColor(79, 70, 229)),
        ("Storage Technologies", "• MongoDB 7.0 Database\n• Local File Evidence Store\n• Vector Similarity Cache\n• Browser LocalStorage", Inches(3.55), Inches(1.05), RGBColor(14, 165, 233)),
        ("Integration & Protocols", "• MCP Server (72 Tools)\n• RESTful JSON HTTP APIs\n• JWT Token Authentication\n• Role-Based Access (RBAC)", Inches(4.75), Inches(1.05), RGBColor(2, 132, 199))
    ]
    for stitle, stext, y_pos, card_h, col in specs:
        card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, y_pos, left_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(248, 250, 252)
        card.line.color.rgb = RGBColor(203, 213, 225)
        pill = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x + Inches(0.12), y_pos - Inches(0.1), left_w - Inches(0.24), Inches(0.26))
        pill.fill.solid()
        pill.fill.fore_color.rgb = col
        pill.line.fill.background()
        p = pill.text_frame.paragraphs[0]
        p.text = stitle
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.2)
        tf.margin_left = Inches(0.1)
        p = tf.paragraphs[0]
        p.text = stext
        p.font.size = Pt(7.8)
        p.font.color.rgb = RGBColor(30, 41, 59)

    # Bottom Yellow Box
    yellow = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, Inches(5.95), Inches(4.3), Inches(1.18))
    yellow.fill.solid()
    yellow.fill.fore_color.rgb = RGBColor(254, 240, 138)
    yellow.line.color.rgb = RGBColor(234, 179, 8)
    tf_y = yellow.text_frame
    tf_y.word_wrap = True
    tf_y.margin_left = Inches(0.12)
    tf_y.margin_top = Inches(0.06)
    phases = [
        "• Scope: Ingest civic complaints, match R&D with CSR funds, track SLAs",
        "• Design: 3-tier microservice architecture (React + Spring Boot + FastAPI + MCP)",
        "• Development: Sentence-Transformers NLP, ELA photo authenticity & MCDA matching",
        "• Deployment: Containerized microservices with MongoDB & JWT security",
        "• Maintenance: Real-time SLA tracking, multi-district analytics & audit logs"
    ]
    for i, phase in enumerate(phases):
        p = tf_y.paragraphs[0] if i == 0 else tf_y.add_paragraph()
        p.text = phase
        p.font.size = Pt(7.3)
        p.font.color.rgb = RGBColor(113, 63, 18)
        if i == 0:
            p.font.bold = True

    # Center Pipeline
    colA = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.25), Inches(0.95), Inches(1.25), Inches(4.85))
    colA.fill.solid()
    colA.fill.fore_color.rgb = RGBColor(241, 245, 249)
    colA.line.color.rgb = RGBColor(203, 213, 225)
    
    colB = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.6), Inches(0.95), Inches(1.25), Inches(4.85))
    colB.fill.solid()
    colB.fill.fore_color.rgb = RGBColor(254, 243, 199)
    colB.line.color.rgb = RGBColor(252, 211, 77)

    colC = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.95), Inches(0.95), Inches(1.3), Inches(4.85))
    colC.fill.solid()
    colC.fill.fore_color.rgb = RGBColor(240, 253, 250)
    colC.line.color.rgb = RGBColor(153, 246, 228)

    colD = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.35), Inches(0.95), Inches(1.2), Inches(4.85))
    colD.fill.solid()
    colD.fill.fore_color.rgb = RGBColor(241, 245, 249)
    colD.line.color.rgb = RGBColor(203, 213, 225)

    colE = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.65), Inches(0.95), Inches(1.9), Inches(4.85))
    colE.fill.solid()
    colE.fill.fore_color.rgb = RGBColor(238, 242, 255)
    colE.line.color.rgb = RGBColor(199, 210, 254)

    colF_x = Inches(10.65)
    colF_w = Inches(2.25)
    status_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, colF_x, Inches(0.95), colF_w, Inches(1.6))
    status_card.fill.solid()
    status_card.fill.fore_color.rgb = RGBColor(254, 243, 199)
    status_card.line.color.rgb = RGBColor(245, 158, 11)
    tf_s = status_card.text_frame
    p = tf_s.paragraphs[0]
    p.text = "Product Status"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(146, 64, 14)
    p2 = tf_s.add_paragraph()
    p2.text = "100% Core System Built & Verified!\n(54/54 Tests Passed, Production Ready)"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(8.2)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(30, 64, 175)

    sec_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, colF_x, Inches(2.65), colF_w, Inches(1.95))
    sec_box.fill.solid()
    sec_box.fill.fore_color.rgb = RGBColor(245, 243, 255)
    sec_box.line.color.rgb = RGBColor(221, 214, 254)
    p = sec_box.text_frame.paragraphs[0]
    p.text = "🔒 Security & Governance\n• JWT + BCrypt Auth\n• Role-Based Access (RBAC)\n• AES-256 Data Encryption\n• SHA-256 Tamper-Proof Audit\n• Human-In-The-Loop AI Gates"
    p.font.size = Pt(7.4)
    p.font.color.rgb = RGBColor(68, 64, 60)

    # Readiness Badges
    badges = [
        ("7", "TECHNOLOGY\nREADINESS LEVEL", RGBColor(14, 116, 144), Inches(4.75)),
        ("6", "DEPLOYMENT\nREADINESS LEVEL", RGBColor(161, 98, 7), Inches(5.45)),
        ("5", "INVESTMENT\nREADINESS LEVEL", RGBColor(30, 58, 138), Inches(6.15))
    ]
    for num, label, col, y_pos in badges:
        circ = slide1.shapes.add_shape(MSO_SHAPE.OVAL, colF_x + Inches(0.1), y_pos, Inches(0.48), Inches(0.48))
        circ.fill.solid()
        circ.fill.fore_color.rgb = col
        circ.line.fill.background()
        p = circ.text_frame.paragraphs[0]
        p.text = num
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        lbl = slide1.shapes.add_textbox(colF_x + Inches(0.65), y_pos - Inches(0.05), Inches(1.5), Inches(0.55))
        p = lbl.text_frame.paragraphs[0]
        p.text = label
        p.font.size = Pt(6.8)
        p.font.bold = True
        p.font.color.rgb = col

    # Footer
    footer1 = slide1.shapes.add_textbox(Inches(4.5), Inches(7.15), Inches(4.4), Inches(0.3))
    p = footer1.text_frame.paragraphs[0]
    p.text = "@SIH Idea submission - CIVIC CONNECT | 3"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = RGBColor(100, 116, 139)


    # =========================================================================
    # SLIDE 2: READINESS LEVELS & MATURITY MATRIX (TRL, DRL, IRL Deep-Dive)
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    slide2.background.fill.solid()
    slide2.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    # Header
    t2 = slide2.shapes.add_textbox(Inches(3.0), Inches(0.35), Inches(7.333), Inches(0.7))
    p = t2.text_frame.paragraphs[0]
    p.text = "READINESS LEVELS & MATURITY MATRIX"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial Black"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 23, 42)

    sub2 = slide2.shapes.add_textbox(Inches(3.0), Inches(0.9), Inches(7.333), Inches(0.4))
    p = sub2.text_frame.paragraphs[0]
    p.text = "Standard Assessment Framework: Technology (TRL) • Deployment (DRL) • Investment (IRL)"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(10)
    p.font.color.rgb = RGBColor(100, 116, 139)

    # 3 Large Cards for TRL, DRL, IRL
    cards_info = [
        ("TRL: 7 / 9", "TECHNOLOGY READINESS LEVEL", RGBColor(14, 116, 144), Inches(0.6),
         "• Current Status: Prototype demonstrated in operational environment.\n• Verified Features: React 18 frontend + Spring Boot 3.4 API + Python FastAPI AI microservices fully integrated.\n• Evidence: Real-time multilingual NLP (Sentence-Transformers) and ELA image authenticity active.\n• Next Milestone (TRL 8): Live municipal field pilot across 3 Jharkhand districts."),

        ("DRL: 6 / 10", "DEPLOYMENT READINESS LEVEL", RGBColor(161, 98, 7), Inches(4.75),
         "• Current Status: Production & Cloud Deployment Ready.\n• Verified Features: 54/54 Automated Test Suite passed across security, data integrity, and end-to-end lifecycle.\n• Evidence: Dockerized microservices, MongoDB database clustering, and Role-Based Access Control (RBAC).\n• Next Milestone (DRL 7): Production rollout on state government cloud server."),

        ("IRL: 5 / 9", "INVESTMENT READINESS LEVEL", RGBColor(30, 58, 138), Inches(8.9),
         "• Current Status: Validated Product-Market Fit & Financial Model.\n• Verified Features: Multi-Stakeholder co-funding architecture (CSR Grants + University R&D + State Grants).\n• Evidence: Automated cost estimation, milestone escrow tracking, and CSR tax compliance support.\n• Next Milestone (IRL 6): Signing initial MoUs with corporate CSR partners (e.g. Tata Steel, NTPC).")
    ]

    for badge_text, full_title, col, x_pos, content in cards_info:
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos, Inches(1.4), Inches(3.8), Inches(4.5))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(248, 250, 252)
        card.line.color.rgb = col
        card.line.width = Pt(1.5)

        # Header bar on card
        bar = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos + Inches(0.1), Inches(1.5), Inches(3.6), Inches(0.6))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()
        p = bar.text_frame.paragraphs[0]
        p.text = badge_text
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Arial Black"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        p2 = bar.text_frame.add_paragraph()
        p2.text = full_title
        p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(7.5)
        p2.font.bold = True
        p2.font.color.rgb = RGBColor(241, 245, 249)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.8)
        tf.margin_left = Inches(0.15)
        tf.margin_right = Inches(0.15)
        p = tf.paragraphs[0]
        p.text = content
        p.font.size = Pt(8.8)
        p.font.color.rgb = RGBColor(30, 41, 59)

    # Bottom Progression Banner
    prog_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(6.05), Inches(12.1), Inches(0.95))
    prog_box.fill.solid()
    prog_box.fill.fore_color.rgb = RGBColor(254, 243, 199)
    prog_box.line.color.rgb = RGBColor(245, 158, 11)
    prog_box.line.width = Pt(1.5)

    tf_p = prog_box.text_frame
    tf_p.word_wrap = True
    tf_p.margin_top = Inches(0.08)
    tf_p.margin_left = Inches(0.15)
    p = tf_p.paragraphs[0]
    p.text = "🎯 6-Month Scalability Roadmap:"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = RGBColor(146, 64, 14)
    p2 = tf_p.add_paragraph()
    p2.text = "• Phase 1 (Current / Hackathon): TRL 7 / DRL 6 / IRL 5 – Verified Full-Stack Prototype with 54/54 Automated Tests Passed.\n• Phase 2 (Months 1-3): TRL 8 / DRL 7 – Deploy Pilot in Ranchi & Dhanbad Municipal Corporations; Onboard 3 Universities.\n• Phase 3 (Months 4-6): TRL 9 / DRL 8 / IRL 6 – State-wide expansion across all 24 districts of Jharkhand with CSR Partnerships."
    p2.font.size = Pt(8.0)
    p2.font.color.rgb = RGBColor(113, 63, 18)

    # Footer Slide 2
    footer2 = slide2.shapes.add_textbox(Inches(4.5), Inches(7.15), Inches(4.4), Inches(0.3))
    p = footer2.text_frame.paragraphs[0]
    p.text = "@SIH Idea submission - CIVIC CONNECT | 4"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = RGBColor(100, 116, 139)


    # =========================================================================
    # SLIDE 3: VERIFIED TECHNOLOGY STACK
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    slide3.background.fill.solid()
    slide3.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    border3 = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(0.3), Inches(12.533), Inches(6.9))
    border3.fill.background()
    border3.line.color.rgb = RGBColor(203, 213, 225)
    border3.line.width = Pt(1.5)

    t3 = slide3.shapes.add_textbox(Inches(3.5), Inches(0.45), Inches(6.333), Inches(0.7))
    p = t3.text_frame.paragraphs[0]
    p.text = "Technology Stack"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Georgia"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.underline = True
    p.font.color.rgb = RGBColor(24, 43, 73)

    tech_data = [
        ("Frontend", [("React 18", RGBColor(6, 182, 212)), ("Vite", RGBColor(168, 85, 247)), ("Bootstrap 5", RGBColor(147, 51, 234)), ("JavaScript", RGBColor(234, 179, 8)), ("React Router", RGBColor(239, 68, 68)), ("Lucide Icons", RGBColor(59, 130, 246))], Inches(1.3)),
        ("Backend", [("Java 17", RGBColor(234, 88, 12)), ("Spring Boot 3.4", RGBColor(22, 163, 74)), ("Spring Security", RGBColor(34, 197, 94)), ("JWT (JJWT)", RGBColor(219, 39, 119)), ("Maven", RGBColor(185, 28, 28))], Inches(2.7)),
        ("AI / ML Microservice", [("Python 3", RGBColor(30, 64, 175)), ("FastAPI", RGBColor(13, 148, 136)), ("PyTorch", RGBColor(234, 88, 12)), ("Transformers", RGBColor(217, 119, 6)), ("Sentence-Transformers", RGBColor(249, 115, 22)), ("Pillow (ELA)", RGBColor(220, 38, 38)), ("Scikit-Learn", RGBColor(234, 88, 12))], Inches(4.1)),
        ("Database & Engine", [("MongoDB 7.0", RGBColor(22, 163, 74)), ("MCP Server (72 Tools)", RGBColor(79, 70, 229)), ("Uvicorn", RGBColor(13, 148, 136))], Inches(5.5))
    ]

    for label, items, y_pos in tech_data:
        lbl_box = slide3.shapes.add_textbox(Inches(0.6), y_pos, Inches(2.2), Inches(1.0))
        p = lbl_box.text_frame.paragraphs[0]
        p.text = label
        p.font.name = "Arial"
        p.font.size = Pt(12.5)
        p.font.bold = True
        p.font.color.rgb = RGBColor(15, 23, 42)

        x_start = Inches(2.9)
        spacing = Inches(1.55)
        for i, (name, col) in enumerate(items):
            bx = x_start + i * spacing
            card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, y_pos, Inches(1.42), Inches(0.95))
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(248, 250, 252)
            card.line.color.rgb = RGBColor(226, 232, 240)
            card.line.width = Pt(1)

            tf_b = card.text_frame
            tf_b.word_wrap = True
            tf_b.margin_top = Inches(0.12)
            p = tf_b.paragraphs[0]
            p.text = name
            p.alignment = PP_ALIGN.CENTER
            p.font.name = "Arial"
            p.font.size = Pt(9.0)
            p.font.bold = True
            p.font.color.rgb = col

    output_path = "CivicConnect_Technical_Approach_Architecture.pptx"
    prs.save(output_path)
    print(f"Updated complete presentation deck saved to: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    update_full_deck()
